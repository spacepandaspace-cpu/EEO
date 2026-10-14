from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DOWNLOADS = Path(r"C:\Users\EEO\Downloads")
PROJECT = Path(__file__).resolve().parent.parent

TRIAL_FILE = DOWNLOADS / "AI教学风险监测灰度测试开通.xlsx"
CLAIM_FILE = DOWNLOADS / "AI额度活动领取情况_2026-09-20.xlsx"
PAID_FILE = DOWNLOADS / "AI额度售卖情况_2026-09-20.xlsx"
USAGE_FILE = PROJECT / "data" / "backend_usage_snapshot_2026-09-21.tsv"
CLAIM_ALIGNMENT_FILE = PROJECT / "data" / "claim_usage_alignment_2026-09-21.tsv"
PAID_REVIEW_FILE = PROJECT / "data" / "paid_uid_mapping_review_2026-09-21.tsv"
IDENTITY_FILE = PROJECT / "data" / "customer_identity_overrides_2026-09-21.tsv"

# 仅写入已核验或由用户明确确认的关系；一个源单元格可拆成多个 UID。
VERIFIED_PHONE_UIDS = {
    "18940956569": ["92555278"],
    "13689589673、 15889762712": ["72605244", "66237956"],
    "18501692570": ["69284500"],
    "15618637836": ["41248694"],
}

PAID_AMOUNT_OVERRIDES = {"ASI_": 5000, "济南新希望教育科技发展有限公司": 20000}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_values(path: Path) -> list[tuple[Any, ...]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    return list(workbook.active.iter_rows(values_only=True))


def read_usage() -> dict[str, dict[str, str]]:
    with USAGE_FILE.open("r", encoding="utf-8-sig", newline="") as handle:
        return {row["uid"]: row for row in csv.DictReader(handle, delimiter="\t")}


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def usage_summary(uids: set[str], usage: dict[str, dict[str, str]]) -> dict[str, int]:
    rows = [usage[uid] for uid in sorted(uids) if uid in usage]
    positive = lambda row, field: int(row.get(field) or 0) > 0
    return {
        "uids": len(uids),
        "uids_in_snapshot": len(rows),
        "used_ever": sum(positive(row, "total") for row in rows),
        "used_prelaunch": sum(positive(row, "preLaunchLessons") for row in rows),
        "used_campaign": sum(positive(row, "campaignLessons") for row in rows),
        "used_postcampaign": sum(positive(row, "postCampaignLessons") for row in rows),
        "lessons_total": sum(int(row.get("total") or 0) for row in rows),
        "lessons_prelaunch": sum(int(row.get("preLaunchLessons") or 0) for row in rows),
        "lessons_campaign": sum(int(row.get("campaignLessons") or 0) for row in rows),
        "lessons_postcampaign": sum(int(row.get("postCampaignLessons") or 0) for row in rows),
    }


def main() -> None:
    usage = read_usage()
    claim_alignment = read_tsv(CLAIM_ALIGNMENT_FILE)
    paid_review = read_tsv(PAID_REVIEW_FILE)
    identity_rows = read_tsv(IDENTITY_FILE)
    identity_by_uid = {row["uid"]: row for row in identity_rows}

    trial_source_rows = []
    trial_account_rows = []
    for row in read_values(TRIAL_FILE)[1:]:
        raw_id, institution, city, owner, status = (text(value) for value in row[:5])
        if not raw_id and not institution:
            continue
        source = {
            "raw_id": raw_id,
            "institution": institution,
            "city": city,
            "owner": owner,
            "status": status,
        }
        trial_source_rows.append(source)
        uids = [raw_id] if raw_id.isdigit() and len(raw_id) < 11 else VERIFIED_PHONE_UIDS.get(raw_id, [])
        for uid in uids:
            override = identity_by_uid.get(uid, {})
            trial_account_rows.append(
                {
                    **source,
                    "uid": uid,
                    "customer": override.get("customerName") or institution,
                    "analysis_status": override.get("analysisStatus") or "纳入试用分析",
                }
            )

    claim_rows = []
    for row in read_values(CLAIM_FILE)[2:]:
        uid = text(row[1])
        if not uid:
            continue
        claim_rows.append(
            {
                "uid": uid,
                "customer": text(row[2]),
                "eeo_account": text(row[3]),
                "owner": text(row[4]),
                "date": text(row[5]),
            }
        )

    paid_rows = []
    for row in read_values(PAID_FILE)[2:]:
        customer = text(row[1])
        if not customer:
            continue
        paid_rows.append(
            {
                "customer": customer,
                "owner": text(row[2]),
                "order": text(row[3]),
                "date": text(row[4]),
                "raw_amount": int(row[6] or 0),
                "amount": PAID_AMOUNT_OVERRIDES.get(customer, int(row[6] or 0)),
            }
        )

    scoped_trial_rows = [
        row
        for row in trial_account_rows
        if row["status"] == "已开通" and row["analysis_status"] != "本期排除"
    ]
    trial_uids = {row["uid"] for row in scoped_trial_rows}
    claim_uids = {row["uid"] for row in claim_rows}
    paid_uids = {row["proposedUid"] for row in paid_review if row["proposedUid"]}
    all_source_uids = trial_uids | claim_uids | paid_uids
    unresolved_trial = [
        row
        for row in trial_source_rows
        if not (
            (row["raw_id"].isdigit() and len(row["raw_id"]) < 11)
            or row["raw_id"] in VERIFIED_PHONE_UIDS
        )
    ]

    claim_name_to_uids: dict[str, set[str]] = {}
    for row in claim_rows:
        claim_name_to_uids.setdefault(row["customer"], set()).add(row["uid"])
    exact_paid_claim = {
        row["customer"]: sorted(claim_name_to_uids[row["customer"]])
        for row in paid_rows
        if row["customer"] in claim_name_to_uids
    }

    output = {
        "snapshot_rows": len(usage),
        "trial": {
            "source_rows": len(trial_source_rows),
            "status_opened_source_rows": sum(
                row["status"] == "已开通" for row in trial_source_rows
            ),
            "status_blank_source_rows": sum(not row["status"] for row in trial_source_rows),
            "scoped_uid_rows_after_split": len(scoped_trial_rows),
            "unique_resolved_uids": len(trial_uids),
            "unresolved": unresolved_trial,
            "usage": usage_summary(trial_uids, usage),
        },
        "claim": {
            "source_rows": len(claim_rows),
            "unique_uids": len(claim_uids),
            "unique_customer_names": len({row["customer"] for row in claim_rows}),
            "usage": usage_summary(claim_uids, usage),
            "event_alignment": {
                "used_on_or_after_record": sum(
                    int(row["lessonsOnOrAfterRecord"]) > 0 for row in claim_alignment
                ),
                "used_within_7_days": sum(
                    int(row["lessonsWithin7Days"]) > 0 for row in claim_alignment
                ),
                "first_used_on_or_after_record": sum(
                    int(row["lessonsBeforeRecord"]) == 0
                    and int(row["lessonsOnOrAfterRecord"]) > 0
                    for row in claim_alignment
                ),
                "used_before_but_not_after": sum(
                    int(row["lessonsBeforeRecord"]) > 0
                    and int(row["lessonsOnOrAfterRecord"]) == 0
                    for row in claim_alignment
                ),
                "lessons_on_or_after_record": sum(
                    int(row["lessonsOnOrAfterRecord"]) for row in claim_alignment
                ),
                "lessons_within_7_days": sum(
                    int(row["lessonsWithin7Days"]) for row in claim_alignment
                ),
            },
        },
        "trial_claim_overlap_uids": len(trial_uids & claim_uids),
        "union_uids": len(trial_uids | claim_uids),
        "all_source_union_uids": len(all_source_uids),
        "all_source_usage": usage_summary(all_source_uids, usage),
        "source_membership": {
            "trial_only": len(trial_uids - claim_uids - paid_uids),
            "claim_only": len(claim_uids - trial_uids - paid_uids),
            "paid_only": len(paid_uids - trial_uids - claim_uids),
            "trial_and_claim_not_paid": len((trial_uids & claim_uids) - paid_uids),
            "trial_and_paid_not_claim": len((trial_uids & paid_uids) - claim_uids),
            "claim_and_paid_not_trial": len((claim_uids & paid_uids) - trial_uids),
            "trial_claim_paid": len(trial_uids & claim_uids & paid_uids),
        },
        "paid": {
            "source_rows": len(paid_rows),
            "raw_detail_amount_sum": sum(row["raw_amount"] for row in paid_rows),
            "corrected_detail_amount_sum": sum(row["amount"] for row in paid_rows),
            "source_total_row_amount": 111000,
            "remaining_difference_vs_total_row": sum(row["amount"] for row in paid_rows)
            - 111000,
            "exact_claim_name_matches": exact_paid_claim,
            "exact_claim_name_match_count": len(exact_paid_claim),
            "mapping_review": {
                "auto_verified": sum(
                    row["matchStatus"] == "自动核验通过" for row in paid_review
                ),
                "pending_confirmation": sum(
                    row["matchStatus"] == "待用户确认" for row in paid_review
                ),
                "user_confirmed": sum(
                    row["matchStatus"] == "用户确认" for row in paid_review
                ),
                "no_candidate": sum(
                    row["matchStatus"] == "未找到候选" for row in paid_review
                ),
                "auto_verified_amount": sum(
                    int(row["amount"])
                    for row in paid_review
                    if row["matchStatus"] == "自动核验通过"
                ),
                "auto_verified_used_after_payment": sum(
                    row["matchStatus"] == "自动核验通过"
                    and int(row["lessonsOnOrAfterPayment"] or 0) > 0
                    for row in paid_review
                ),
                "auto_verified_lessons_after_payment": sum(
                    int(row["lessonsOnOrAfterPayment"] or 0)
                    for row in paid_review
                    if row["matchStatus"] == "自动核验通过"
                ),
                "all_mapped_used_after_payment": sum(
                    int(row["lessonsOnOrAfterPayment"] or 0) > 0 for row in paid_review
                ),
                "all_mapped_lessons_after_payment": sum(
                    int(row["lessonsOnOrAfterPayment"] or 0) for row in paid_review
                ),
            },
        },
    }
    print(json.dumps(output, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
