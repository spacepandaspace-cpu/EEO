from __future__ import annotations

import argparse
import csv
import json
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DOWNLOADS = Path(r"C:\Users\EEO\Downloads")
PROJECT = Path(__file__).resolve().parent.parent

TRIAL_FILE = DOWNLOADS / "AI教学风险监测灰度测试开通.xlsx"
CLAIM_FILE = DOWNLOADS / "AI额度活动领取情况_2026-09-20.xlsx"
PAID_FILE = DOWNLOADS / "AI额度售卖情况_2026-09-20.xlsx"
IDENTITY_FILE = PROJECT / "data" / "customer_identity_overrides_2026-09-21.tsv"
PAID_REVIEW_FILE = PROJECT / "data" / "paid_uid_mapping_review_2026-09-21.tsv"
USAGE_FILE = PROJECT / "data" / "backend_usage_snapshot_2026-09-21.tsv"

PHONE_UIDS = {
    "18940956569": ["92555278"],
    "13689589673、 15889762712": ["72605244", "66237956"],
    "18501692570": ["69284500"],
    "15618637836": ["41248694"],
}

PAID_AMOUNT_OVERRIDES = {"ASI_": 5000, "济南新希望教育科技发展有限公司": 20000}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def rows(path: Path) -> list[tuple[Any, ...]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    return list(workbook.active.iter_rows(values_only=True))


def tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def consultant_parts(raw: str) -> tuple[str, str]:
    if "-EEO" in raw:
        name, identifier = raw.rsplit("-", 1)
        return name, identifier
    return raw, ""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--write", action="store_true")
    args = parser.parse_args()
    identity = {row["uid"]: row for row in tsv(IDENTITY_FILE)}
    paid_map = {row["paidCustomer"]: row for row in tsv(PAID_REVIEW_FILE)}
    usage = {row["uid"]: row for row in tsv(USAGE_FILE)}

    evidence: dict[str, list[dict[str, str]]] = defaultdict(list)

    for row in rows(TRIAL_FILE)[1:]:
        raw_id, institution, city, owner, status = (text(value) for value in row[:5])
        if not raw_id and not institution:
            continue
        if status != "已开通":
            continue
        uids = [raw_id] if raw_id.isdigit() and len(raw_id) < 11 else PHONE_UIDS.get(raw_id, [])
        for uid in uids:
            evidence[uid].append(
                {
                    "source": "trial",
                    "date": "",
                    "customer": institution,
                    "eeo_account": "",
                    "consultant": owner,
                    "city": city,
                }
            )

    for row in rows(CLAIM_FILE)[2:]:
        uid = text(row[1])
        if not uid:
            continue
        evidence[uid].append(
            {
                "source": "claim",
                "date": text(row[5]).replace("/", "-"),
                "customer": text(row[2]),
                "eeo_account": text(row[3]),
                "consultant": text(row[4]),
                "city": "",
            }
        )

    for row in rows(PAID_FILE)[2:]:
        paid_customer = text(row[1])
        if not paid_customer:
            continue
        mapping = paid_map[paid_customer]
        uid = mapping["proposedUid"]
        evidence[uid].append(
            {
                "source": "paid",
                "date": text(row[4]).replace("/", "-"),
                "customer": mapping["customerLabelFromSource"] or paid_customer,
                "eeo_account": "",
                "consultant": text(row[2]),
                "city": "",
            }
        )

    paid_details: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows(PAID_FILE)[2:]:
        customer = text(row[1])
        if not customer:
            continue
        uid = paid_map[customer]["proposedUid"]
        paid_details[uid].append(
            {
                "date": text(row[4]).replace("/", "-"),
                "amount": PAID_AMOUNT_OVERRIDES.get(customer, int(row[6] or 0)),
                "order": text(row[3]),
            }
        )

    master = []
    for uid in sorted(evidence, key=lambda value: int(value)):
        items = evidence[uid]
        override = identity.get(uid, {})
        claims = [item for item in items if item["source"] == "claim"]
        paid = [item for item in items if item["source"] == "paid"]
        trials = [item for item in items if item["source"] == "trial"]

        selected = (paid or claims or trials)[-1]
        customer = override.get("customerName") or selected["customer"]
        eeo_account = override.get("eeoAccountName") or next(
            (item["eeo_account"] for item in claims if item["eeo_account"]), ""
        )
        eeo_blank_confirmed = (
            not eeo_account and "EEO账号名从缺" in (override.get("note") or "")
        )

        dated_owner_items = [item for item in items if item["consultant"]]
        dated_owner_items.sort(key=lambda item: (item["date"], item["source"] == "paid"))
        consultant_raw = dated_owner_items[-1]["consultant"] if dated_owner_items else ""
        consultant_name, consultant_id = consultant_parts(consultant_raw)
        consultant_override = (override.get("currentConsultant") or "").strip()
        if consultant_override:
            consultant_name = consultant_override
            consultant_id = next(
                (
                    consultant_parts(item["consultant"])[1]
                    for item in reversed(dated_owner_items)
                    if consultant_parts(item["consultant"])[0] == consultant_override
                ),
                "",
            )

        customer_names = sorted({item["customer"] for item in items if item["customer"]})
        consultant_names = sorted(
            {consultant_parts(item["consultant"])[0] for item in items if item["consultant"]}
        )
        payments = paid_details.get(uid, [])
        usage_row = usage.get(uid, {})

        master.append(
            {
                "uid": uid,
                "eeoAccountName": eeo_account,
                "eeoAccountNameStatus": "确认从缺" if eeo_blank_confirmed else "已核验",
                "customerName": customer,
                "consultant": consultant_name,
                "consultantId": consultant_id,
                "trial": bool(trials),
                "claim": bool(claims),
                "claimDate": claims[0]["date"] if claims else "",
                "paid": bool(payments),
                "paymentCount": len(payments),
                "paymentTotal": sum(item["amount"] for item in payments),
                "paymentDetails": payments,
                "used": int(usage_row.get("total") or 0) > 0,
                "analyzedLessons": int(usage_row.get("total") or 0),
                "lastUsage": usage_row.get("lastUsage", ""),
                "customerNameEvidence": customer_names,
                "consultantEvidence": consultant_names,
            }
        )

    missing_eeo = [
        row["uid"]
        for row in master
        if not row["eeoAccountName"] and row["eeoAccountNameStatus"] != "确认从缺"
    ]
    customer_conflicts = [
        {
            "uid": row["uid"],
            "selected": row["customerName"],
            "evidence": row["customerNameEvidence"],
        }
        for row in master
        if len(row["customerNameEvidence"]) > 1
        and not (identity.get(row["uid"], {}).get("customerName") or "").strip()
    ]
    consultant_conflicts = [
        {
            "uid": row["uid"],
            "selected": row["consultant"],
            "evidence": row["consultantEvidence"],
        }
        for row in master
        if len(row["consultantEvidence"]) > 1
        and not (identity.get(row["uid"], {}).get("currentConsultant") or "").strip()
    ]

    if args.write:
        customer_conflict_uids = {row["uid"] for row in customer_conflicts}
        consultant_conflict_uids = {row["uid"] for row in consultant_conflicts}
        output_path = PROJECT / "data" / "customer_master_audit_2026-09-21.tsv"
        fields = [
            "uid",
            "eeoAccountName",
            "eeoAccountNameStatus",
            "customerName",
            "consultant",
            "consultantId",
            "trial",
            "claim",
            "claimDate",
            "paid",
            "paymentCount",
            "paymentTotal",
            "paymentDetails",
            "used",
            "analyzedLessons",
            "lastUsage",
            "customerNameEvidence",
            "consultantEvidence",
            "auditStatus",
        ]
        with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t")
            writer.writeheader()
            for row in master:
                issues = []
                if (
                    not row["eeoAccountName"]
                    and row["eeoAccountNameStatus"] != "确认从缺"
                ):
                    issues.append("EEO账号名待确认")
                if row["uid"] in customer_conflict_uids:
                    issues.append("客户标准名待确认")
                if row["uid"] in consultant_conflict_uids:
                    issues.append("当前顾问待确认")
                output = {key: row.get(key, "") for key in fields}
                output["paymentDetails"] = json.dumps(
                    row["paymentDetails"], ensure_ascii=False, separators=(",", ":")
                )
                output["customerNameEvidence"] = json.dumps(
                    row["customerNameEvidence"], ensure_ascii=False, separators=(",", ":")
                )
                output["consultantEvidence"] = json.dumps(
                    row["consultantEvidence"], ensure_ascii=False, separators=(",", ":")
                )
                output["auditStatus"] = "；".join(issues) if issues else "已核验"
                writer.writerow(output)

    print(
        json.dumps(
            {
                "master_rows": len(master),
                "missing_eeo_count": len(missing_eeo),
                "missing_eeo_uids": missing_eeo,
                "customer_conflict_count": len(customer_conflicts),
                "customer_conflicts": customer_conflicts,
                "consultant_conflict_count": len(consultant_conflicts),
                "consultant_conflicts": consultant_conflicts,
                "corrected_payment_total": sum(
                    row["paymentTotal"] for row in master
                ),
                "written": str(PROJECT / "data" / "customer_master_audit_2026-09-21.tsv")
                if args.write
                else "",
            },
            ensure_ascii=True,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
